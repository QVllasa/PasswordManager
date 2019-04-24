from docx import Document
from docx.shared import *

document = Document()

table = document.add_table(rows = 2, cols = 6, style='Table Grid')
a = table.cell(0,1)
b = table.cell(0,5)
A = a.merge(b)
hdr_cells = table.rows[0].cells

hdr_cells[0].text = 'Username'
hdr_cells[0].text.style = 'Intense'
hdr_cells[1].text = 'qendrimvllasa@siemens.com'
psw_cells = table.rows[1].cells
psw_cells[0].text = "Password"
psw_cells[1].text = 'blablabla'
psw_cells[2].text = 'IoT Extension'
psw_cells[3].text = 'blablabla'
psw_cells[4].text = 'Integration'
psw_cells[5].text = 'blablabla'





document.save('test' + '.docx')